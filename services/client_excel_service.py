"""Import/export seguro de clientes mediante archivos XLSX.

La importación trabaja únicamente sobre clientes de la empresa actual y no permite
modificar campos internos, credenciales, tokens ni compañía.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


MAX_IMPORT_ROWS = 10000
EXCEL_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Campos que el usuario puede importar/exportar. company_id, token y timestamps
# quedan fuera deliberadamente por seguridad/integridad.
IMPORT_FIELDS = (
    "id",
    "name",
    "email",
    "phone",
    "whatsapp",
    "address",
    "city",
    "province",
    "postal_code",
    "birthday",
    "balance",
    "credit_limit",
    "notes",
    "observations",
    "email_marketing_consent",
    "whatsapp_marketing_consent",
    "account_current_enabled",
    "active",
)

DISPLAY_HEADERS = {
    "id": "id",
    "name": "nombre",
    "email": "email",
    "phone": "telefono",
    "whatsapp": "whatsapp",
    "address": "direccion",
    "city": "ciudad",
    "province": "provincia",
    "postal_code": "codigo_postal",
    "birthday": "fecha_nacimiento",
    "balance": "saldo",
    "credit_limit": "limite_credito",
    "notes": "notas",
    "observations": "observaciones",
    "email_marketing_consent": "consentimiento_email",
    "whatsapp_marketing_consent": "consentimiento_whatsapp",
    "account_current_enabled": "cuenta_corriente_habilitada",
    "active": "activo",
}

ALIASES = {
    "id": {"id", "cliente id", "id cliente", "codigo cliente", "codigo"},
    "name": {"name", "nombre", "cliente", "razon social", "razon_social", "nombre cliente", "nombre completo"},
    "email": {"email", "correo", "correo electronico", "correo electrónico", "mail"},
    "phone": {"phone", "telefono", "tel", "telefono fijo", "celular", "movil", "móvil"},
    "whatsapp": {"whatsapp", "whatsapp telefono", "whatsapp tel", "wa"},
    "address": {"address", "direccion", "dirección", "domicilio", "calle"},
    "city": {"city", "ciudad", "localidad"},
    "province": {"province", "provincia", "estado"},
    "postal_code": {"postal code", "postal_code", "codigo postal", "código postal", "cp"},
    "birthday": {"birthday", "cumpleanos", "cumpleaños", "fecha nacimiento", "fecha de nacimiento", "nacimiento"},
    "balance": {"balance", "saldo", "cuenta corriente", "deuda", "saldo actual"},
    "credit_limit": {"credit limit", "credit_limit", "limite credito", "límite crédito", "limite de credito", "límite de crédito", "credito maximo"},
    "notes": {"notes", "notas", "nota"},
    "observations": {"observations", "observaciones", "observacion", "observación"},
    "email_marketing_consent": {"email marketing consent", "email_marketing_consent", "consentimiento email", "consentimiento de email", "marketing email"},
    "whatsapp_marketing_consent": {"whatsapp marketing consent", "whatsapp_marketing_consent", "consentimiento whatsapp", "consentimiento de whatsapp", "marketing whatsapp"},
    "account_current_enabled": {"account current enabled", "account_current_enabled", "cuenta corriente habilitada", "cuenta corriente activa", "cc habilitada"},
    "active": {"active", "activo", "estado", "habilitado", "habilitada"},
}

CONSENT_VALUES = {
    "unknown": "unknown",
    "desconocido": "unknown",
    "sin definir": "unknown",
    "no definido": "unknown",
    "pendiente": "unknown",
    "opted_in": "opted_in",
    "opt in": "opted_in",
    "si": "opted_in",
    "sí": "opted_in",
    "acepta": "opted_in",
    "aceptado": "opted_in",
    "true": "opted_in",
    "1": "opted_in",
    "yes": "opted_in",
    "opted_out": "opted_out",
    "opt out": "opted_out",
    "no": "opted_out",
    "rechaza": "opted_out",
    "rechazado": "opted_out",
    "false": "opted_out",
    "0": "opted_out",
}

TRUE_VALUES = {"1", "true", "verdadero", "si", "sí", "yes", "y", "activo", "activa", "habilitado", "habilitada", "on"}
FALSE_VALUES = {"0", "false", "falso", "no", "n", "inactivo", "inactiva", "deshabilitado", "deshabilitada", "off"}


class ClientImportError(ValueError):
    pass


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def normalize_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", normalize_header(value))


def normalize_text(value: Any) -> str:
    return str(value or "").strip()


def parse_numeric(value: Any, field_label: str, *, default: Decimal | None = None) -> Decimal | None:
    if value in (None, ""):
        return default
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    normalized = normalize_text(value)
    normalized = re.sub(r"(?i)(ars|usd|eur|mxn|brl|\$|€|£)", "", normalized)
    normalized = re.sub(r"[^0-9,.-]", "", normalized.replace(" ", ""))
    if not normalized:
        return default
    if "," in normalized and "." in normalized:
        if normalized.rfind(",") > normalized.rfind("."):
            normalized = normalized.replace(".", "").replace(",", ".")
        else:
            normalized = normalized.replace(",", "")
    elif "," in normalized:
        normalized = normalized.replace(",", ".")
    try:
        return Decimal(normalized)
    except (InvalidOperation, ValueError) as exc:
        raise ClientImportError(f"Valor numérico inválido en {field_label}: {value!r}") from exc


def parse_bool(value: Any, field_label: str, *, default: bool = False) -> bool:
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    normalized = normalize_header(value)
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    raise ClientImportError(f"Valor booleano inválido en {field_label}: {value!r}. Usá Sí/No.")


def parse_consent(value: Any, field_label: str, *, default: str = "unknown") -> str:
    if value in (None, ""):
        return default
    normalized = normalize_header(value)
    if normalized in CONSENT_VALUES:
        return CONSENT_VALUES[normalized]
    raise ClientImportError(
        f"Consentimiento inválido en {field_label}: {value!r}. Usá unknown, opted_in u opted_out."
    )


def parse_date(value: Any, field_label: str) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # openpyxl deja seriales numéricos si el Excel perdió formato de fecha.
        from openpyxl.utils.datetime import from_excel
        try:
            return from_excel(value).date() if isinstance(from_excel(value), datetime) else from_excel(value)
        except Exception as exc:
            raise ClientImportError(f"Fecha inválida en {field_label}: {value!r}") from exc
    text_value = normalize_text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text_value, fmt).date()
        except ValueError:
            continue
    raise ClientImportError(
        f"Fecha inválida en {field_label}: {value!r}. Usá DD/MM/AAAA o AAAA-MM-DD."
    )


def classify_headers(raw_headers: tuple[Any, ...]) -> dict[str, int]:
    normalized_aliases = {
        field: {normalize_key(alias) for alias in aliases}
        for field, aliases in ALIASES.items()
    }
    positions: dict[str, int] = {}
    for idx, raw in enumerate(raw_headers):
        compact = normalize_key(raw)
        if not compact:
            continue
        for field, accepted in normalized_aliases.items():
            if compact in accepted and field not in positions:
                positions[field] = idx
                break
    return positions


def find_header_row(workbook) -> tuple[Any, int, dict[str, int]]:
    for sheet in workbook.worksheets:
        for row_number, raw_row in enumerate(
            sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1
        ):
            positions = classify_headers(raw_row)
            if "name" in positions:
                return sheet, row_number, positions
    raise ClientImportError(
        "No se encontró una fila de encabezados válida. Debe existir al menos la columna Nombre."
    )


def _cell(row: tuple[Any, ...], positions: dict[str, int], field: str, default: Any = None) -> Any:
    index = positions.get(field)
    if index is None or index >= len(row):
        return default
    value = row[index]
    return default if value is None else value


def build_records(upload) -> list[dict[str, Any]]:
    if not upload or not upload.filename:
        raise ClientImportError("Seleccioná un archivo Excel .xlsx.")
    if not upload.filename.lower().endswith(".xlsx"):
        raise ClientImportError("El archivo debe ser .xlsx.")

    workbook = load_workbook(upload, read_only=True, data_only=True)
    try:
        sheet, header_row, positions = find_header_row(workbook)
        records: list[dict[str, Any]] = []
        identity_keys: dict[str, int] = {}

        for row_number, raw_row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if row_number > MAX_IMPORT_ROWS + header_row:
                raise ClientImportError(f"El archivo supera el límite de {MAX_IMPORT_ROWS} filas.")

            raw_values = list(raw_row)
            if not any(value not in (None, "") for value in raw_values):
                continue

            name = normalize_text(_cell(raw_row, positions, "name"))
            if not name:
                raise ClientImportError(f"Fila {row_number}: el nombre es obligatorio.")

            raw_id = _cell(raw_row, positions, "id")
            client_id = None
            if raw_id not in (None, ""):
                try:
                    client_id = int(float(raw_id))
                except (TypeError, ValueError) as exc:
                    raise ClientImportError(f"Fila {row_number}: ID de cliente inválido: {raw_id!r}") from exc
                if client_id <= 0:
                    raise ClientImportError(f"Fila {row_number}: ID de cliente inválido: {raw_id!r}")

            email = normalize_text(_cell(raw_row, positions, "email")) or None
            phone = normalize_text(_cell(raw_row, positions, "phone")) or None
            whatsapp = normalize_text(_cell(raw_row, positions, "whatsapp")) or None

            # Una fila duplicada dentro del mismo archivo se rechaza en vez de
            # sobrescribir silenciosamente una fila anterior.
            identity = None
            if client_id:
                identity = f"id:{client_id}"
            elif email:
                identity = f"email:{normalize_key(email)}"
            elif whatsapp:
                identity = f"whatsapp:{normalize_key(whatsapp)}"
            elif phone:
                identity = f"phone:{normalize_key(phone)}"
            if identity:
                if identity in identity_keys:
                    raise ClientImportError(
                        f"Fila {row_number}: cliente duplicado dentro del archivo; "
                        f"ya aparece en la fila {identity_keys[identity]}."
                    )
                identity_keys[identity] = row_number

            records.append(
                {
                    "row_number": row_number,
                    "id": client_id,
                    "name": name,
                    "email": email,
                    "phone": phone,
                    "whatsapp": whatsapp,
                    "address": normalize_text(_cell(raw_row, positions, "address")) or None,
                    "city": normalize_text(_cell(raw_row, positions, "city")) or None,
                    "province": normalize_text(_cell(raw_row, positions, "province")) or None,
                    "postal_code": normalize_text(_cell(raw_row, positions, "postal_code")) or None,
                    "birthday": parse_date(_cell(raw_row, positions, "birthday"), "fecha_nacimiento"),
                    "balance": parse_numeric(_cell(raw_row, positions, "balance"), "saldo", default=Decimal("0.00")),
                    "credit_limit": parse_numeric(_cell(raw_row, positions, "credit_limit"), "limite_credito", default=Decimal("0.00")),
                    "notes": normalize_text(_cell(raw_row, positions, "notes")) or None,
                    "observations": normalize_text(_cell(raw_row, positions, "observations")) or None,
                    "email_marketing_consent": parse_consent(_cell(raw_row, positions, "email_marketing_consent"), "consentimiento_email"),
                    "whatsapp_marketing_consent": parse_consent(_cell(raw_row, positions, "whatsapp_marketing_consent"), "consentimiento_whatsapp"),
                    "account_current_enabled": parse_bool(_cell(raw_row, positions, "account_current_enabled"), "cuenta_corriente_habilitada", default=False),
                    "active": parse_bool(_cell(raw_row, positions, "active"), "activo", default=True),
                }
            )
        if not records:
            raise ClientImportError("No se encontraron clientes para importar.")
        return records
    finally:
        workbook.close()


def _build_identity_lookup(Client, scope_query, records):
    ids = {r["id"] for r in records if r["id"]}
    emails = {r["email"] for r in records if r["email"]}
    phones = {r["phone"] for r in records if r["phone"]}
    whatsapps = {r["whatsapp"] for r in records if r["whatsapp"]}

    clients = scope_query(Client.query, Client).filter(Client.id.in_(ids)).all() if ids else []
    clients += scope_query(Client.query, Client).filter(Client.email.in_(emails)).all() if emails else []
    clients += scope_query(Client.query, Client).filter(Client.phone.in_(phones)).all() if phones else []
    clients += scope_query(Client.query, Client).filter(Client.whatsapp.in_(whatsapps)).all() if whatsapps else []

    by_id = {c.id: c for c in clients}
    by_email = {normalize_key(c.email): c for c in clients if c.email}
    by_phone = {normalize_key(c.phone): c for c in clients if c.phone}
    by_whatsapp = {normalize_key(c.whatsapp): c for c in clients if c.whatsapp}
    return by_id, by_email, by_phone, by_whatsapp


def resolve_existing(record, lookups):
    by_id, by_email, by_phone, by_whatsapp = lookups
    if record["id"] and record["id"] in by_id:
        return by_id[record["id"]]
    if record["email"]:
        match = by_email.get(normalize_key(record["email"]))
        if match:
            return match
    if record["whatsapp"]:
        match = by_whatsapp.get(normalize_key(record["whatsapp"]))
        if match:
            return match
    if record["phone"]:
        match = by_phone.get(normalize_key(record["phone"]))
        if match:
            return match
    return None


def apply_record(client, record):
    for field in IMPORT_FIELDS:
        if field == "id":
            continue
        setattr(client, field, record[field])


def create_workbook(*, clients, client_stats, include_guide: bool = True) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clientes"

    headers = [
        DISPLAY_HEADERS[field]
        for field in IMPORT_FIELDS
    ] + [
        "compras",
        "total_comprado",
        "presupuestos",
        "total_presupuestado",
    ]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for client in clients:
        stats = client_stats.get(client.id, {})
        row = [
            client.id,
            client.name,
            client.email or "",
            client.phone or "",
            client.whatsapp or "",
            client.address or "",
            client.city or "",
            client.province or "",
            client.postal_code or "",
            client.birthday.isoformat() if client.birthday else "",
            float(client.balance or 0),
            float(client.credit_limit or 0),
            client.notes or "",
            client.observations or "",
            client.email_marketing_consent or "unknown",
            client.whatsapp_marketing_consent or "unknown",
            "Sí" if client.account_current_enabled else "No",
            "Sí" if client.active else "No",
            stats.get("purchase_count", 0),
            float(stats.get("total_spent", 0)),
            stats.get("quote_count", 0),
            float(stats.get("total_quoted", 0)),
        ]
        sheet.append(row)

    widths = {
        1: 10, 2: 28, 3: 30, 4: 20, 5: 20, 6: 34, 7: 20, 8: 20, 9: 14,
        10: 16, 11: 14, 12: 16, 13: 34, 14: 34, 15: 24, 16: 26, 17: 24, 18: 12,
        19: 12, 20: 18, 21: 14, 22: 20,
    }
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    if include_guide:
        guide = workbook.create_sheet("Guía")
        guide_rows = [
            ("Importación de clientes", "Usá la hoja Clientes para cargar nuevos clientes o actualizar existentes."),
            ("Identificación", "Para actualizar con precisión, conservá el id exportado. Sin id se intenta coincidir por email, WhatsApp o teléfono."),
            ("Campos obligatorios", "nombre es obligatorio."),
            ("Fechas", "fecha_nacimiento: DD/MM/AAAA o AAAA-MM-DD."),
            ("Valores Sí/No", "cuenta_corriente_habilitada y activo aceptan Sí/No, 1/0, true/false."),
            ("Consentimientos", "consentimiento_email y consentimiento_whatsapp aceptan unknown, opted_in u opted_out."),
            ("Columnas informativas", "compras, total_comprado, presupuestos y total_presupuestado son informativas: el sistema las calcula y no las importa."),
            ("Seguridad", "No se exportan/importan company_id, token de baja de marketing, created_at ni updated_at."),
            ("Límite", f"Máximo {MAX_IMPORT_ROWS} filas por importación."),
        ]
        for row in guide_rows:
            guide.append(row)
        guide.column_dimensions["A"].width = 28
        guide.column_dimensions["B"].width = 105
        for cell in guide[1]:
            cell.font = Font(bold=True)
        guide["A1"] = "Tema"
        guide["B1"] = "Detalle"
        guide.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
