"""Modulo de compras: proveedores, ordenes e ingreso de mercaderia."""

from collections import defaultdict
from datetime import datetime
from io import BytesIO

from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from sqlalchemy import false, or_
from stockarmobile.helpers.numbers import safe_float

bp = Blueprint("purchases", __name__)


def _to_float(value, default=0.0):
    return safe_float(value, default)


def _can_manage_purchases():
    return getattr(current_user, "role", None) in {"admin", "superadmin"}


def _resolve_company_id(required=False):
    from app import Company
    role = getattr(current_user, "role", None)
    if role == "superadmin":
        company_id = request.values.get("company_id", type=int)
    else:
        company_id = getattr(current_user, "company_id", None)
    if required and not company_id:
        return None
    if company_id and Company.query.filter_by(id=company_id).first() is None:
        return None
    return company_id


def _company_scope(query, model, company_id):
    if not hasattr(model, "company_id"):
        return query
    if company_id is None:
        if getattr(current_user, "role", None) != "superadmin":
            return query.filter(false())
        return query
    return query.filter(model.company_id == company_id)


def _require_company_context_or_forbid(company_id):
    if getattr(current_user, "role", None) == "superadmin":
        return
    if company_id is None:
        abort(403)


def _superadmin_companies():
    from app import Company
    return Company.query.filter(Company.active.is_(True)).order_by(Company.name.asc()).all()


def _purchase_access_guard():
    if not _can_manage_purchases():
        abort(403)
    return None


def _form_list(name):
    """Accepts the new [] form fields and the legacy single-value fields."""
    values = request.form.getlist(f"{name}[]")
    if values:
        return values
    values = request.form.getlist(name)
    if values:
        return values
    value = request.form.get(name)
    return [value] if value is not None else []


def _parse_purchase_lines():
    product_values = _form_list("product_id")
    quantity_values = _form_list("quantity")
    cost_values = _form_list("unit_cost")
    lengths = {len(product_values), len(quantity_values), len(cost_values)}
    if lengths != {len(product_values)}:
        raise ValueError("Completa producto, cantidad y costo en todas las líneas.")

    lines = []
    for index, (raw_product, raw_quantity, raw_cost) in enumerate(zip(product_values, quantity_values, cost_values), start=1):
        product_id = int(raw_product) if str(raw_product or "").strip().isdigit() else None
        quantity = _to_float(raw_quantity)
        unit_cost = _to_float(raw_cost)
        if not product_id or quantity <= 0 or unit_cost < 0:
            raise ValueError(f"La línea {index} tiene producto, cantidad o costo inválido.")
        lines.append({"product_id": product_id, "quantity": quantity, "unit_cost": unit_cost})
    if not lines:
        raise ValueError("Agregá al menos un producto a la compra.")
    return lines


def _apply_product_purchase_totals(product_totals):
    """Apply stock and moving-average cost for aggregated purchase quantities."""
    from app import Product

    for product, purchase_quantity, purchase_value in product_totals.values():
        previous_stock = float(product.stock or 0)
        previous_cost = float(product.cost_price or 0)
        new_stock = previous_stock + purchase_quantity
        product.stock = new_stock
        product.cost_price = (
            ((previous_stock * previous_cost) + purchase_value) / new_stock
            if new_stock
            else 0
        )
        product.margin = float(product.price or 0) - float(product.cost_price or 0)
        product.profit_percent = (
            product.margin / float(product.cost_price) * 100
            if float(product.cost_price or 0)
            else 0
        )


@bp.route("/", methods=["GET", "POST"])
@login_required
def index():
    from app import Product, PurchaseItem, PurchaseOrder, Supplier, db, record_audit, utcnow
    _purchase_access_guard()
    company_id = _resolve_company_id(required=False)
    _require_company_context_or_forbid(company_id)
    if getattr(current_user, "role", None) == "superadmin" and company_id is None:
        return render_template("compras/index.html", suppliers=[], products=[], purchases=[], selected_company_id=None, companies=_superadmin_companies(), search="", status_filter="active", supplier_stats={}, company_required=True)
    if request.method == "POST":
        if company_id is None:
            flash("Selecciona una empresa para continuar.", "warning")
            return redirect(url_for("purchases.index"))
        supplier_id = request.form.get("supplier_id", type=int)
        supplier = None
        if supplier_id:
            supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id, Supplier.active.is_(True)).first()
            if supplier is None:
                flash("Proveedor no encontrado para esta empresa.", "danger")
                return redirect(url_for("purchases.index", company_id=company_id))
        try:
            lines = _parse_purchase_lines()
            product_totals = {}
            resolved_lines = []
            for line in lines:
                product = _company_scope(db.session.query(Product), Product, company_id).filter(Product.id == line["product_id"], Product.active.is_(True)).first()
                if product is None:
                    raise ValueError(f"Producto #{line['product_id']} no encontrado para esta empresa.")
                resolved_lines.append((product, line))
                if product.id not in product_totals:
                    product_totals[product.id] = [product, 0.0, 0.0]
                product_totals[product.id][1] += line["quantity"]
                product_totals[product.id][2] += line["quantity"] * line["unit_cost"]
            subtotal = sum(line["quantity"] * line["unit_cost"] for _, line in resolved_lines)
            order = PurchaseOrder(
                supplier_id=supplier.id if supplier else None,
                company_id=company_id,
                date=utcnow(),
                status="recibida",
                subtotal=subtotal,
                total_amount=subtotal,
                note=(request.form.get("note") or "").strip() or None,
            )
            db.session.add(order)
            db.session.flush()
            for product, line in resolved_lines:
                db.session.add(PurchaseItem(purchase_order_id=order.id, product_id=product.id, quantity=line["quantity"], unit_cost=line["unit_cost"]))
            _apply_product_purchase_totals(product_totals)
            record_audit(
                action="purchase_create",
                entity="purchase_order",
                entity_id=order.id,
                detail=f"Compra registrada productos={len(resolved_lines)} proveedores={supplier.id if supplier else 'sin proveedor'} total={subtotal:.2f}",
            )
            db.session.commit()
            flash(f"Compra #{order.id} registrada con {len(resolved_lines)} producto(s) y stock actualizado.", "success")
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "danger")
        except Exception:
            db.session.rollback()
            flash("No se pudo registrar la compra. Revisá los datos e intentá nuevamente.", "danger")
        return redirect(url_for("purchases.index", company_id=company_id))
    suppliers = _company_scope(Supplier.query.filter_by(active=True), Supplier, company_id).order_by(Supplier.name).all()
    products = _company_scope(Product.query.filter_by(active=True), Product, company_id).order_by(Product.name).all()
    purchases = _company_scope(PurchaseOrder.query, PurchaseOrder, company_id).order_by(PurchaseOrder.date.desc()).limit(20).all()
    return render_template("compras/index.html", suppliers=suppliers, products=products, purchases=purchases, selected_company_id=company_id, companies=_superadmin_companies() if getattr(current_user, "role", None) == "superadmin" else [], search="", status_filter="active", supplier_stats={}, company_required=False)


@bp.route("/proveedores", methods=["GET", "POST"])
@login_required
def suppliers_panel():
    from app import PurchaseOrder, Supplier, db, record_audit
    _purchase_access_guard()
    company_id = _resolve_company_id(required=False)
    _require_company_context_or_forbid(company_id)
    if getattr(current_user, "role", None) == "superadmin" and company_id is None:
        return render_template("compras/suppliers.html", suppliers=[], supplier_stats={}, selected_company_id=None, companies=_superadmin_companies(), search="", status_filter="active", company_required=True)
    if request.method == "POST":
        if company_id is None:
            flash("Selecciona una empresa para continuar.", "warning")
            return redirect(url_for("purchases.suppliers_panel"))
        name = (request.form.get("name") or "").strip()
        if not name:
            flash("El proveedor necesita nombre.", "danger")
            return redirect(url_for("purchases.suppliers_panel", company_id=company_id))
        supplier = Supplier(company_id=company_id, name=name, email=(request.form.get("email") or "").strip() or None, phone=(request.form.get("phone") or "").strip() or None, whatsapp=(request.form.get("whatsapp") or "").strip() or None, address=(request.form.get("address") or "").strip() or None, notes=(request.form.get("notes") or "").strip() or None, active=True)
        db.session.add(supplier)
        db.session.flush()
        record_audit(action="supplier_create", entity="supplier", entity_id=supplier.id, detail=f"Proveedor creado: {supplier.name}")
        db.session.commit()
        flash("Proveedor creado.", "success")
        return redirect(url_for("purchases.suppliers_panel", company_id=company_id))
    search = (request.args.get("q") or "").strip()
    status_filter = (request.args.get("status") or "active").strip().lower()
    supplier_query = _company_scope(Supplier.query, Supplier, company_id)
    if search:
        like = f"%{search}%"
        supplier_query = supplier_query.filter(or_(Supplier.name.ilike(like), Supplier.email.ilike(like), Supplier.phone.ilike(like), Supplier.whatsapp.ilike(like), Supplier.address.ilike(like), Supplier.notes.ilike(like)))
    if status_filter == "active":
        supplier_query = supplier_query.filter(Supplier.active.is_(True))
    elif status_filter == "inactive":
        supplier_query = supplier_query.filter(Supplier.active.is_(False))
    suppliers = supplier_query.order_by(Supplier.name.asc(), Supplier.id.desc()).all()
    stats_rows = _company_scope(db.session.query(PurchaseOrder.supplier_id, db.func.count(PurchaseOrder.id).label("purchase_count"), db.func.coalesce(db.func.sum(PurchaseOrder.total_amount), 0).label("total_amount")), PurchaseOrder, company_id).filter(PurchaseOrder.supplier_id.isnot(None)).group_by(PurchaseOrder.supplier_id).all()
    supplier_stats = {int(row.supplier_id): {"purchase_count": int(row.purchase_count or 0), "total_amount": float(row.total_amount or 0)} for row in stats_rows if row.supplier_id is not None}
    return render_template("compras/suppliers.html", suppliers=suppliers, supplier_stats=supplier_stats, selected_company_id=company_id, companies=_superadmin_companies() if getattr(current_user, "role", None) == "superadmin" else [], search=search, status_filter=status_filter, company_required=False)


@bp.route("/proveedores/<int:supplier_id>/compras")
@login_required
def supplier_purchases(supplier_id):
    from app import PurchaseOrder, Supplier, db
    _purchase_access_guard()
    company_id = _resolve_company_id(required=True)
    _require_company_context_or_forbid(company_id)
    supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id).first_or_404()
    purchases = _filtered_supplier_purchases(PurchaseOrder, supplier, company_id, db)
    search = (request.args.get("q") or "").strip()
    status = (request.args.get("status") or "").strip().lower()
    date_from = (request.args.get("date_from") or "").strip()
    date_to = (request.args.get("date_to") or "").strip()
    return render_template("compras/supplier_purchases.html", supplier=supplier, purchases=purchases, selected_company_id=company_id, search=search, status=status, date_from=date_from, date_to=date_to)


@bp.route("/proveedores/<int:supplier_id>/compras/export.xlsx")
@login_required
def supplier_purchases_excel(supplier_id):
    from app import PurchaseItem, PurchaseOrder, Supplier, db
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from sqlalchemy.orm import selectinload

    _purchase_access_guard()
    company_id = _resolve_company_id(required=True)
    supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id).first_or_404()
    query = _supplier_purchases_query(PurchaseOrder, supplier, company_id, db)
    purchases = query.options(selectinload(PurchaseOrder.items).selectinload(PurchaseItem.product)).all()

    wb = Workbook(); ws = wb.active; ws.title = "Compras"
    headers = ["Compra", "Fecha", "Proveedor", "Producto", "Cantidad", "Costo unitario", "Subtotal línea", "Total compra", "Estado", "Observaciones"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for purchase in purchases:
        items = list(purchase.items or [])
        if not items:
            ws.append([purchase.id, purchase.date.strftime("%Y-%m-%d %H:%M") if purchase.date else "", supplier.name, "", "", "", "", float(purchase.total_amount or 0), purchase.status or "", purchase.note or ""])
            continue
        for item in items:
            quantity = float(item.quantity or 0)
            unit_cost = float(item.unit_cost or 0)
            ws.append([
                purchase.id,
                purchase.date.strftime("%Y-%m-%d %H:%M") if purchase.date else "",
                supplier.name,
                item.product.name if item.product else f"Producto {item.product_id}",
                quantity,
                unit_cost,
                quantity * unit_cost,
                float(purchase.total_amount or 0),
                purchase.status or "",
                purchase.note or "",
            ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=8):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'
    for col in ws.columns:
        width = max(len(str(cell.value or "")) for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 42)
    if ws.max_row == 1:
        ws.append(["", "", "", "No hay compras que coincidan con los filtros.", "", "", "", "", "", ""])

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"compras_{supplier.id}.xlsx")


@bp.route("/proveedores/<int:supplier_id>/compras/export.pdf")
@login_required
def supplier_purchases_pdf(supplier_id):
    from app import Supplier, PurchaseOrder, db
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    _purchase_access_guard()
    company_id = _resolve_company_id(required=True)
    supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id).first_or_404()
    purchases = _filtered_supplier_purchases(PurchaseOrder, supplier, company_id, db)
    buf = BytesIO(); pdf = canvas.Canvas(buf, pagesize=A4); width, height = A4
    pdf.setTitle(f"Compras - {supplier.name}"); y = height - 45
    pdf.setFont("Helvetica-Bold", 15); pdf.drawString(40, y, "StockArmobile - Compras por proveedor"); y -= 22
    pdf.setFont("Helvetica", 10); pdf.drawString(40, y, f"Proveedor: {supplier.name}"); y -= 18
    pdf.drawString(40, y, f"Generado: {datetime.utcnow().strftime('%Y-%m-%d %H:%M')}"); y -= 24
    pdf.setFont("Helvetica-Bold", 9); pdf.drawString(40, y, "Compra"); pdf.drawString(90, y, "Fecha"); pdf.drawString(185, y, "Productos"); pdf.drawRightString(width - 40, y, "Total"); y -= 15
    pdf.setFont("Helvetica", 8)
    for purchase in purchases:
        if y < 45: pdf.showPage(); y = height - 45; pdf.setFont("Helvetica", 8)
        names = ", ".join((item.product.name if item.product else f"Producto {item.product_id}") for item in purchase.items) or "Sin detalle"
        if len(names) > 65: names = names[:62] + "..."
        pdf.drawString(40, y, f"#{purchase.id}"); pdf.drawString(90, y, purchase.date.strftime("%d/%m/%Y") if purchase.date else ""); pdf.drawString(185, y, names); pdf.drawRightString(width - 40, y, f"${float(purchase.total_amount or 0):,.2f}"); y -= 13
    pdf.save(); buf.seek(0)
    return send_file(buf, mimetype="application/pdf", as_attachment=True, download_name=f"compras_{supplier.id}.pdf")


def _supplier_purchases_query(PurchaseOrder, supplier, company_id, db):
    query = _company_scope(PurchaseOrder.query, PurchaseOrder, company_id).filter(PurchaseOrder.supplier_id == supplier.id)
    search = (request.args.get("q") or "").strip(); status = (request.args.get("status") or "").strip().lower()
    if search:
        from app import Product, PurchaseItem
        like = f"%{search}%"
        query = (
            query.outerjoin(PurchaseItem, PurchaseOrder.id == PurchaseItem.purchase_order_id)
            .outerjoin(Product, PurchaseItem.product_id == Product.id)
            .filter(or_(PurchaseOrder.note.ilike(like), db.cast(PurchaseOrder.id, db.String).ilike(like), Product.name.ilike(like)))
            .distinct()
        )
    if status: query = query.filter(PurchaseOrder.status.ilike(status))
    for key, op in (("date_from", ">=",), ("date_to", "<")):
        value = (request.args.get(key) or "").strip()
        if value:
            try:
                parsed = datetime.strptime(value, "%Y-%m-%d")
                if key == "date_to": parsed = parsed.replace(hour=23, minute=59, second=59)
                query = query.filter(PurchaseOrder.date >= parsed if op == ">=" else PurchaseOrder.date < parsed)
            except ValueError:
                pass
    return query.order_by(PurchaseOrder.date.desc(), PurchaseOrder.id.desc())


def _filtered_supplier_purchases(PurchaseOrder, supplier, company_id, db):
    return _supplier_purchases_query(PurchaseOrder, supplier, company_id, db).all()


@bp.route("/proveedores/<int:supplier_id>/compras/<int:purchase_id>/editar", methods=["GET", "POST"])
@login_required
def edit_supplier_purchase(supplier_id, purchase_id):
    from app import Product, PurchaseItem, PurchaseOrder, Supplier, db, record_audit
    _purchase_access_guard(); company_id = _resolve_company_id(required=True); _require_company_context_or_forbid(company_id)
    supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id).first_or_404()
    purchase = _company_scope(db.session.query(PurchaseOrder), PurchaseOrder, company_id).filter(PurchaseOrder.id == purchase_id, PurchaseOrder.supplier_id == supplier.id).first_or_404()
    products = _company_scope(Product.query.filter_by(active=True), Product, company_id).order_by(Product.name).all()
    if request.method == "POST":
        try:
            lines = _parse_purchase_lines()
            resolved_lines = []
            new_totals = {}
            for line in lines:
                product = _company_scope(db.session.query(Product), Product, company_id).filter(Product.id == line["product_id"], Product.active.is_(True)).first()
                if product is None:
                    raise ValueError(f"Producto #{line['product_id']} no encontrado para esta empresa.")
                resolved_lines.append((product, line))
                if product.id not in new_totals:
                    new_totals[product.id] = [product, 0.0, 0.0]
                new_totals[product.id][1] += line["quantity"]
                new_totals[product.id][2] += line["quantity"] * line["unit_cost"]

            old_totals = {}
            for item in list(purchase.items or []):
                product = _company_scope(db.session.query(Product), Product, company_id).filter(Product.id == item.product_id).first()
                if product is None:
                    raise ValueError(f"El producto #{item.product_id} asociado a la compra ya no existe.")
                if product.id not in old_totals:
                    old_totals[product.id] = [product, 0.0, 0.0]
                old_totals[product.id][1] += float(item.quantity or 0)
                old_totals[product.id][2] += float(item.quantity or 0) * float(item.unit_cost or 0)

            affected_ids = set(old_totals) | set(new_totals)
            for product_id in affected_ids:
                old_qty = old_totals.get(product_id, [None, 0.0, 0.0])[1]
                new_qty = new_totals.get(product_id, [None, 0.0, 0.0])[1]
                new_value = new_totals.get(product_id, [None, 0.0, 0.0])[2]
                product = (new_totals.get(product_id) or old_totals.get(product_id))[0]
                current_stock = float(product.stock or 0)
                base_stock = current_stock - old_qty
                if base_stock < -1e-9:
                    raise ValueError(f"No se puede editar {product.name}: su stock actual es menor al stock aportado por esta compra.")
                base_value = max(base_stock, 0.0) * float(product.cost_price or 0)
                updated_stock = base_stock + new_qty
                if updated_stock < -1e-9:
                    raise ValueError(f"La edición dejaría stock negativo para {product.name}.")
                product.stock = updated_stock
                product.cost_price = ((base_value + new_value) / updated_stock) if updated_stock > 0 else 0
                product.margin = float(product.price or 0) - float(product.cost_price or 0)
                product.profit_percent = product.margin / float(product.cost_price) * 100 if float(product.cost_price or 0) else 0

            for old_item in list(purchase.items or []):
                db.session.delete(old_item)
            for product, line in resolved_lines:
                db.session.add(PurchaseItem(purchase_order_id=purchase.id, product_id=product.id, quantity=line["quantity"], unit_cost=line["unit_cost"]))
            subtotal = sum(line["quantity"] * line["unit_cost"] for _, line in resolved_lines)
            purchase.supplier_id = supplier.id
            purchase.status = (request.form.get("status") or purchase.status or "recibida").strip()
            purchase.note = (request.form.get("note") or "").strip() or None
            purchase.subtotal = subtotal
            purchase.total_amount = subtotal
            record_audit(action="purchase_update", entity="purchase_order", entity_id=purchase.id, detail=f"Compra actualizada productos={len(resolved_lines)} proveedor={supplier.id} total={subtotal:.2f}")
            db.session.commit()
            flash(f"Compra #{purchase.id} actualizada correctamente con {len(resolved_lines)} producto(s).", "success")
            return redirect(url_for("purchases.supplier_purchases", supplier_id=supplier.id, company_id=company_id))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "danger")
            return redirect(url_for("purchases.edit_supplier_purchase", supplier_id=supplier.id, purchase_id=purchase.id, company_id=company_id))
        except Exception:
            db.session.rollback()
            flash("No se pudo actualizar la compra. No se modificó el stock.", "danger")
            return redirect(url_for("purchases.edit_supplier_purchase", supplier_id=supplier.id, purchase_id=purchase.id, company_id=company_id))
    return render_template("compras/edit_purchase.html", supplier=supplier, purchase=purchase, products=products, selected_company_id=company_id)


@bp.route("/proveedores/<int:supplier_id>/compras/import.xlsx", methods=["POST"])
@login_required
def import_supplier_purchases(supplier_id):
    from app import Product, PurchaseItem, PurchaseOrder, Supplier, db, record_audit, utcnow
    from openpyxl import load_workbook
    _purchase_access_guard(); company_id = _resolve_company_id(required=True); _require_company_context_or_forbid(company_id)
    supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id).first_or_404()
    upload = request.files.get("file")
    if not upload or not upload.filename.lower().endswith((".xlsx", ".xlsm")):
        flash("Selecciona un archivo Excel .xlsx válido.", "danger"); return redirect(url_for("purchases.supplier_purchases", supplier_id=supplier.id, company_id=company_id))
    try:
        wb = load_workbook(upload, read_only=True, data_only=True); ws = wb.active; rows = list(ws.iter_rows(values_only=True)); wb.close()
        if not rows: raise ValueError("El Excel está vacío.")
        headers = {str(v or "").strip().lower(): i for i, v in enumerate(rows[0])}
        def col(*names):
            for name in names:
                if name in headers: return headers[name]
            return None
        pcol = col("producto", "product", "codigo", "código", "id producto")
        qcol = col("cantidad", "quantity")
        ccol = col("costo unitario", "unit_cost", "costo", "precio costo")
        if qcol is None or ccol is None or pcol is None: raise ValueError("Faltan columnas obligatorias: Producto, Cantidad y Costo unitario.")
        valid_lines = []
        totals = {}
        for row in rows[1:]:
            if not row or all(v in (None, "") for v in row):
                continue
            raw_product = row[pcol]; quantity = _to_float(row[qcol]); unit_cost = _to_float(row[ccol])
            product = None
            if str(raw_product).strip().isdigit():
                product = _company_scope(db.session.query(Product), Product, company_id).filter(Product.id == int(raw_product), Product.active.is_(True)).first()
            if product is None:
                product = _company_scope(Product.query, Product, company_id).filter(Product.name.ilike(str(raw_product).strip()), Product.active.is_(True)).first()
            if product is None or quantity <= 0 or unit_cost < 0:
                continue
            valid_lines.append((product, quantity, unit_cost))
            if product.id not in totals:
                totals[product.id] = [product, 0.0, 0.0]
            totals[product.id][1] += quantity
            totals[product.id][2] += quantity * unit_cost
        if not valid_lines: raise ValueError("No se encontraron filas válidas para importar.")
        subtotal = sum(quantity * unit_cost for _, quantity, unit_cost in valid_lines)
        order = PurchaseOrder(supplier_id=supplier.id, company_id=company_id, date=utcnow(), status="recibida", subtotal=subtotal, total_amount=subtotal, note="Importada desde Excel")
        db.session.add(order); db.session.flush()
        for product, quantity, unit_cost in valid_lines:
            db.session.add(PurchaseItem(purchase_order_id=order.id, product_id=product.id, quantity=quantity, unit_cost=unit_cost))
        _apply_product_purchase_totals(totals)
        record_audit(action="purchase_import_excel", entity="supplier", entity_id=supplier.id, detail=f"Compra importada desde Excel id={order.id} productos={len(valid_lines)} total={subtotal:.2f}")
        db.session.commit(); flash(f"Se importó la compra #{order.id} con {len(valid_lines)} producto(s).", "success")
    except ValueError as exc:
        db.session.rollback(); flash(str(exc), "danger")
    except Exception:
        db.session.rollback(); flash("No se pudo importar el Excel. Revisá el formato y los productos.", "danger")
    return redirect(url_for("purchases.supplier_purchases", supplier_id=supplier.id, company_id=company_id))


@bp.route("/proveedores/<int:supplier_id>/compras/plantilla.xlsx")
@login_required
def supplier_purchase_template(supplier_id):
    from app import Supplier, db
    from openpyxl import Workbook
    _purchase_access_guard(); company_id = _resolve_company_id(required=True)
    supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id,).first_or_404()
    wb = Workbook(); ws = wb.active; ws.title = "Compras"; ws.append(["Producto", "Cantidad", "Costo unitario"]); ws.append(["Ejemplo", 1, 1000])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"plantilla_compras_{supplier.id}.xlsx")


@bp.route("/proveedores/<int:supplier_id>/update", methods=["POST"])
@login_required
def update_supplier(supplier_id):
    from app import Supplier, db, record_audit
    _purchase_access_guard(); company_id = _resolve_company_id(required=True)
    if company_id is None: flash("Empresa inválida.", "danger"); return redirect(url_for("purchases.suppliers_panel"))
    supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id).first_or_404()
    name = (request.form.get("name") or "").strip()
    if not name: flash("El proveedor necesita nombre.", "danger"); return redirect(url_for("purchases.suppliers_panel", company_id=company_id))
    supplier.name = name; supplier.email = (request.form.get("email") or "").strip() or None; supplier.phone = (request.form.get("phone") or "").strip() or None; supplier.whatsapp = (request.form.get("whatsapp") or "").strip() or None; supplier.address = (request.form.get("address") or "").strip() or None; supplier.notes = (request.form.get("notes") or "").strip() or None
    db.session.add(supplier); record_audit(action="supplier_update", entity="supplier", entity_id=supplier.id, detail=f"Proveedor actualizado: {supplier.name}"); db.session.commit(); flash("Proveedor actualizado.", "success")
    return redirect(url_for("purchases.suppliers_panel", company_id=company_id))


@bp.route("/proveedores/<int:supplier_id>/toggle", methods=["POST"])
@login_required
def toggle_supplier(supplier_id):
    from app import Supplier, db, record_audit
    _purchase_access_guard(); company_id = _resolve_company_id(required=True)
    if company_id is None: flash("Empresa inválida.", "danger"); return redirect(url_for("purchases.suppliers_panel"))
    supplier = _company_scope(db.session.query(Supplier), Supplier, company_id).filter(Supplier.id == supplier_id).first_or_404()
    supplier.active = not bool(supplier.active); db.session.add(supplier); action = "supplier_activate" if supplier.active else "supplier_deactivate"; detail = "Proveedor reactivado" if supplier.active else "Proveedor desactivado"
    record_audit(action=action, entity="supplier", entity_id=supplier.id, detail=f"{detail}: {supplier.name}"); db.session.commit(); flash("Proveedor actualizado.", "success")
    return redirect(url_for("purchases.suppliers_panel", company_id=company_id))
