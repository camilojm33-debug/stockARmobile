from io import BytesIO
from types import SimpleNamespace

import pytest
from werkzeug.datastructures import FileStorage

from services.client_excel_service import (
    ClientImportError,
    build_records,
    create_workbook,
)


def _upload(rows, headers):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clientes"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return FileStorage(stream=buffer, filename="clientes.xlsx")


def test_client_excel_import_parses_full_business_fields():
    upload = _upload(
        [[
            12,
            "Ana Pérez",
            "ana@example.com",
            "3794112233",
            "5493794112233",
            "Av. Siempre Viva 123",
            "Resistencia",
            "Chaco",
            "3500",
            "07/10/1985",
            "1.234,50",
            "5000",
            "Cliente mayorista",
            "Compra mensual",
            "opted_in",
            "No",
            "Sí",
        ]],
        [
            "ID",
            "Nombre",
            "Email",
            "Teléfono",
            "WhatsApp",
            "Dirección",
            "Ciudad",
            "Provincia",
            "Código Postal",
            "Fecha de Nacimiento",
            "Saldo",
            "Límite de Crédito",
            "Notas",
            "Observaciones",
            "Consentimiento Email",
            "Consentimiento WhatsApp",
            "Cuenta Corriente Habilitada",
        ],
    )

    record = build_records(upload)[0]

    assert record["id"] == 12
    assert record["name"] == "Ana Pérez"
    assert record["province"] == "Chaco"
    assert record["postal_code"] == "3500"
    assert str(record["balance"]) == "1234.50"
    assert str(record["credit_limit"]) == "5000"
    assert record["birthday"].isoformat() == "1985-10-07"
    assert record["email_marketing_consent"] == "opted_in"
    assert record["whatsapp_marketing_consent"] == "unknown"
    assert record["account_current_enabled"] is True


def test_client_excel_import_rejects_duplicate_identity_rows():
    upload = _upload(
        [
            ["Cliente A", "a@example.com"],
            ["Cliente B", "a@example.com"],
        ],
        ["Nombre", "Email"],
    )

    with pytest.raises(ClientImportError, match="duplicado"):
        build_records(upload)


def test_client_excel_import_requires_name():
    upload = _upload(
        [["", "a@example.com"]],
        ["Nombre", "Email"],
    )

    with pytest.raises(ClientImportError, match="nombre es obligatorio"):
        build_records(upload)


def test_client_excel_export_contains_importable_and_derived_columns():
    client = SimpleNamespace(
        id=7,
        name="Cliente Demo",
        email="demo@example.com",
        phone="123",
        whatsapp="456",
        address="Calle 1",
        city="Resistencia",
        province="Chaco",
        postal_code="3500",
        birthday=None,
        balance=100,
        credit_limit=500,
        notes="Notas",
        observations="Obs",
        email_marketing_consent="unknown",
        whatsapp_marketing_consent="opted_in",
        account_current_enabled=True,
        active=True,
    )

    buffer = create_workbook(
        clients=[client],
        client_stats={
            7: {
                "purchase_count": 4,
                "total_spent": 12000.5,
                "quote_count": 2,
                "total_quoted": 18000,
            }
        },
    )

    from openpyxl import load_workbook

    workbook = load_workbook(buffer, read_only=True, data_only=True)
    sheet = workbook["Clientes"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]

    assert "nombre" in headers
    assert "provincia" in headers
    assert "saldo" in headers
    assert "compras" in headers
    assert "total_comprado" in headers
    assert "presupuestos" in headers
    assert row[headers.index("nombre")] == "Cliente Demo"
    assert row[headers.index("compras")] == 4
    assert row[headers.index("total_comprado")] == 12000.5
    assert "company_id" not in headers
    assert "marketing_unsubscribe_token" not in headers
    workbook.close()


def test_client_excel_template_is_exposed_in_client_screen():
    from pathlib import Path

    html = Path("templates/clientes/index.html").read_text(encoding="utf-8")
    assert "Exportar Excel" in html
    assert "Importar Excel" in html
    assert "Descargar plantilla" in html
    assert 'enctype="multipart/form-data"' in html
