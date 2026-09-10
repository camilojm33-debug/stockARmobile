import io

from openpyxl import Workbook, load_workbook

from app import Company, Product, PurchaseItem, PurchaseOrder, Supplier, User, db


def _login(client, user):
    with client.session_transaction() as session:
        session["_user_id"] = str(user.id)
        session["_fresh"] = True


def _setup_purchase_context():
    company = Company(name="Empresa Compras Test", active=True)
    db.session.add(company)
    db.session.flush()
    user = User(
        username="compras_admin",
        email="compras-admin@test.local",
        password_hash="test-password-hash",
        role="admin",
        active=True,
        company_id=company.id,
    )
    supplier = Supplier(company_id=company.id, name="Proveedor Test", active=True)
    product_a = Product(
        barcode="CMP-A",
        name="Producto A",
        price=150,
        cost_price=100,
        stock=0,
        min_stock=1,
        active=True,
        company_id=company.id,
    )
    product_b = Product(
        barcode="CMP-B",
        name="Producto B",
        price=300,
        cost_price=200,
        stock=0,
        min_stock=1,
        active=True,
        company_id=company.id,
    )
    db.session.add_all([user, supplier, product_a, product_b])
    db.session.commit()
    return company, user, supplier, product_a, product_b


def test_create_purchase_with_multiple_products_updates_one_order_and_stock(app):
    client = app.test_client()
    with app.app_context():
        company, user, supplier, product_a, product_b = _setup_purchase_context()
        _login(client, user)

        response = client.post(
            "/compras/",
            data={
                "company_id": str(company.id),
                "supplier_id": str(supplier.id),
                "product_id[]": [str(product_a.id), str(product_b.id)],
                "quantity[]": ["2", "3"],
                "unit_cost[]": ["25", "50"],
                "note": "Compra multiproducto",
            },
            follow_redirects=False,
        )
        assert response.status_code in {302, 303}

        purchase = PurchaseOrder.query.filter_by(company_id=company.id, supplier_id=supplier.id).one()
        assert len(purchase.items) == 2
        assert float(purchase.total_amount) == 200.0
        assert float(db.session.get(Product, product_a.id).stock) == 2.0
        assert float(db.session.get(Product, product_b.id).stock) == 3.0


def test_supplier_purchase_excel_contains_multiple_items(app):
    client = app.test_client()
    with app.app_context():
        company, user, supplier, product_a, product_b = _setup_purchase_context()
        purchase = PurchaseOrder(
            supplier_id=supplier.id,
            company_id=company.id,
            status="recibida",
            subtotal=200,
            total_amount=200,
        )
        db.session.add(purchase)
        db.session.flush()
        db.session.add_all([
            PurchaseItem(purchase_order_id=purchase.id, product_id=product_a.id, quantity=2, unit_cost=25),
            PurchaseItem(purchase_order_id=purchase.id, product_id=product_b.id, quantity=3, unit_cost=50),
        ])
        db.session.commit()
        _login(client, user)

        response = client.get(f"/compras/proveedores/{supplier.id}/compras/export.xlsx?company_id={company.id}")
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
        rows = list(workbook["Compras"].iter_rows(values_only=True))
        workbook.close()

        assert rows[0][:4] == ("Compra", "Fecha", "Proveedor", "Producto")
        products = {row[3] for row in rows[1:] if row[3]}
        assert products == {"Producto A", "Producto B"}
        assert any(row[4] == 2.0 and row[5] == 25.0 for row in rows[1:])
        assert any(row[4] == 3.0 and row[5] == 50.0 for row in rows[1:])


def test_edit_purchase_with_multiple_products_replaces_lines_and_stock(app):
    client = app.test_client()
    with app.app_context():
        company, user, supplier, product_a, product_b = _setup_purchase_context()
        purchase = PurchaseOrder(
            supplier_id=supplier.id,
            company_id=company.id,
            status="recibida",
            subtotal=200,
            total_amount=200,
        )
        db.session.add(purchase)
        db.session.flush()
        db.session.add_all([
            PurchaseItem(purchase_order_id=purchase.id, product_id=product_a.id, quantity=2, unit_cost=25),
            PurchaseItem(purchase_order_id=purchase.id, product_id=product_b.id, quantity=3, unit_cost=50),
        ])
        product_a.stock = 2
        product_b.stock = 3
        db.session.commit()
        _login(client, user)

        response = client.post(
            f"/compras/proveedores/{supplier.id}/compras/{purchase.id}/editar",
            data={
                "company_id": str(company.id),
                "product_id[]": [str(product_a.id), str(product_b.id)],
                "quantity[]": ["4", "1"],
                "unit_cost[]": ["30", "55"],
                "status": "recibida",
                "note": "Editada",
            },
            follow_redirects=False,
        )
        assert response.status_code in {302, 303}

        refreshed = db.session.get(PurchaseOrder, purchase.id)
        assert len(refreshed.items) == 2
        assert float(refreshed.total_amount) == 175.0
        assert float(db.session.get(Product, product_a.id).stock) == 4.0
        assert float(db.session.get(Product, product_b.id).stock) == 1.0


def test_import_excel_creates_one_multiproduct_purchase(app):
    client = app.test_client()
    with app.app_context():
        company, user, supplier, product_a, product_b = _setup_purchase_context()
        _login(client, user)

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Producto", "Cantidad", "Costo unitario"])
        sheet.append([product_a.name, 2, 25])
        sheet.append([product_b.name, 3, 50])
        payload = io.BytesIO()
        workbook.save(payload)
        payload.seek(0)

        response = client.post(
            f"/compras/proveedores/{supplier.id}/compras/import.xlsx?company_id={company.id}",
            data={"file": (payload, "compras.xlsx")},
            content_type="multipart/form-data",
            follow_redirects=False,
        )
        assert response.status_code in {302, 303}
        purchase = PurchaseOrder.query.filter_by(company_id=company.id, supplier_id=supplier.id).one()
        assert len(purchase.items) == 2
        assert float(purchase.total_amount) == 200.0
