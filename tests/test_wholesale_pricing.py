from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest
from flask_login import login_user, logout_user

from app import Client, Company, Plan, Product, Subscription, User, db, utcnow
from services.plan_service import PlanService
from services.wholesale_pricing_service import (
    apply_price_snapshot,
    apply_wholesale_prices,
    product_price_catalog,
    resolve_product_price,
)
from wholesale_pricing import PriceList, PriceListItem


@pytest.fixture
def wholesale_database(app):
    company = Company(
        name="Empresa mayorista QA",
        active=True,
        trial_ends_at=utcnow() + timedelta(days=10),
    )
    other_company = Company(
        name="Otra empresa QA",
        active=True,
        trial_ends_at=utcnow() + timedelta(days=10),
    )
    db.session.add_all([company, other_company])
    db.session.flush()

    PlanService.ensure_defaults(db.session)
    business_plan = Plan.query.filter_by(code="business").first()
    entrepreneur_plan = Plan.query.filter_by(code="entrepreneur").first()
    assert business_plan is not None and entrepreneur_plan is not None
    business_tokens = [item.strip() for item in str(business_plan.features_json or "").replace(";", ",").split(",") if item.strip()]
    if "all" in {item.lower() for item in business_tokens}:
        business_tokens = ["inventario", "ventas", "clientes", "compras", "caja", "reportes", "excel", "kardex"]
    # The test app omits these unrelated blueprints; keep their menu links locked
    # so the wholesale page can render without requiring unrelated route modules.
    business_tokens = [
        item for item in business_tokens
        if item.lower() not in {"pricing_controller", "promotions", "wholesale_pricing"}
    ]
    business_tokens.append("wholesale_pricing")
    business_plan.features_json = ",".join(business_tokens)
    entrepreneur_tokens = [item.strip() for item in str(entrepreneur_plan.features_json or "").replace(";", ",").split(",") if item.strip()]
    entrepreneur_tokens = [item for item in entrepreneur_tokens if item.lower() != "wholesale_pricing"]
    if "all" in {item.lower() for item in entrepreneur_tokens}:
        entrepreneur_tokens = ["inventario", "ventas", "clientes", "reportes", "excel"]
    entrepreneur_plan.features_json = ",".join(entrepreneur_tokens)
    subscription = Subscription(company_id=company.id, plan_id=business_plan.id, status="active")
    db.session.add(subscription)
    db.session.flush()

    admin = User(
        username="wholesale_admin",
        email="wholesale-admin@test.local",
        password_hash="test-password-hash",
        role="admin",
        active=True,
        company_id=company.id,
    )
    employee = User(
        username="wholesale_employee",
        email="wholesale-employee@test.local",
        password_hash="test-password-hash",
        role="user",
        active=True,
        company_id=company.id,
    )
    db.session.add_all([admin, employee])
    db.session.flush()

    product = Product(
        company_id=company.id,
        barcode="WHOLESALE-001",
        name="Producto mayorista QA",
        category="Bebidas",
        price=Decimal("100.00"),
        cost_price=Decimal("50.00"),
        discount=Decimal("5.00"),
        stock=100,
        min_stock=1,
        active=True,
        sale_type="unidad",
        unit_measure="u",
    )
    other_product = Product(
        company_id=other_company.id,
        barcode="OTHER-001",
        name="Producto de otra empresa",
        price=Decimal("300.00"),
        cost_price=Decimal("100.00"),
        discount=Decimal("0.00"),
        stock=10,
        min_stock=1,
        active=True,
        sale_type="unidad",
        unit_measure="u",
    )
    db.session.add_all([product, other_product])
    db.session.flush()

    price_list = PriceList(
        company_id=company.id,
        name="Mayorista QA",
        code="mayorista-qa",
        active=True,
        created_by_user_id=admin.id,
    )
    db.session.add(price_list)
    db.session.flush()
    db.session.add_all([
        PriceListItem(
            company_id=company.id,
            price_list_id=price_list.id,
            product_id=product.id,
            min_quantity=Decimal("1.000"),
            unit_price=Decimal("80.00"),
            active=True,
        ),
        PriceListItem(
            company_id=company.id,
            price_list_id=price_list.id,
            product_id=product.id,
            min_quantity=Decimal("12.000"),
            unit_price=Decimal("70.00"),
            active=True,
        ),
    ])
    wholesale_client = Client(
        company_id=company.id,
        name="Cliente mayorista QA",
        active=True,
        customer_type="mayorista",
        price_list_id=price_list.id,
    )
    retail_client = Client(
        company_id=company.id,
        name="Cliente minorista QA",
        active=True,
        customer_type="minorista",
        price_list_id=price_list.id,
    )
    db.session.add_all([wholesale_client, retail_client])
    db.session.commit()
    return {
        "company": company,
        "other_company": other_company,
        "business_plan": business_plan,
        "entrepreneur_plan": entrepreneur_plan,
        "subscription": subscription,
        "admin": admin,
        "employee": employee,
        "product": product,
        "other_product": other_product,
        "price_list": price_list,
        "wholesale_client": wholesale_client,
        "retail_client": retail_client,
    }


def test_resolver_uses_the_highest_quantity_tier_and_retail_fallback(wholesale_database):
    data = wholesale_database
    company_id = data["company"].id
    product = data["product"]
    wholesale_client = data["wholesale_client"]
    retail_client = data["retail_client"]

    one = resolve_product_price(
        company_id=company_id, product=product, quantity=1, client=wholesale_client
    )
    dozen = resolve_product_price(
        company_id=company_id, product=product, quantity=12, client=wholesale_client
    )
    retail = resolve_product_price(
        company_id=company_id, product=product, quantity=12, client=retail_client
    )
    wrong_tenant = resolve_product_price(
        company_id=company_id, product=product, quantity=12,
        client=Client(company_id=data["other_company"].id, customer_type="mayorista", price_list_id=data["price_list"].id),
    )

    assert one["source"] == "mayorista"
    assert one["price"] == Decimal("80.00")
    assert dozen["source"] == "mayorista"
    assert dozen["price"] == Decimal("70.00")
    assert dozen["price_list_id"] == data["price_list"].id
    assert retail["source"] == "minorista"
    assert retail["price"] == Decimal("100.00")
    assert wrong_tenant["source"] == "minorista"
    assert wrong_tenant["price"] == Decimal("100.00")


def test_apply_wholesale_price_does_not_mutate_product_or_stack_legacy_discount(wholesale_database):
    data = wholesale_database
    product = data["product"]
    original_price = Decimal(str(product.price))
    lines = [{
        "product": product,
        "quantity": Decimal("12"),
        "price": Decimal(str(product.price)),
        "discount": Decimal("60.00"),
    }]

    applied = apply_wholesale_prices(
        lines, company_id=data["company"].id, client=data["wholesale_client"]
    )

    assert applied[0]["price"] == Decimal("70.00")
    assert applied[0]["discount"] == Decimal("0.00")
    assert applied[0]["price_source"] == "mayorista"
    assert applied[0]["price_list_id"] == data["price_list"].id
    assert product.price == original_price


def test_product_price_catalog_is_company_scoped_and_returns_quantity_tiers(wholesale_database):
    data = wholesale_database
    catalog = product_price_catalog(
        company_id=data["company"].id,
        product_ids=[data["product"].id, data["other_product"].id],
    )

    assert data["product"].id in catalog
    tiers = catalog[data["product"].id][str(data["price_list"].id)]
    assert tiers == [
        {"min_quantity": 1.0, "price": 80.0},
        {"min_quantity": 12.0, "price": 70.0},
    ]
    assert data["other_product"].id not in catalog


def test_quote_parser_applies_assigned_list_and_removes_legacy_discount(wholesale_database, app):
    import quotes

    data = wholesale_database
    with app.test_request_context("/presupuestos/nuevo"):
        login_user(data["employee"])
        try:
            items = quotes._parse_items(
                {"items": [{
                    "product_id": data["product"].id,
                    "description": data["product"].name,
                    "quantity": 12,
                    "unit_price": 100,
                    "discount": 60,
                }]},
                client=data["wholesale_client"],
            )
        finally:
            logout_user()

    assert len(items) == 1
    assert items[0]["unit_price"] == Decimal("70.00")
    assert items[0]["discount"] == Decimal("0.00")
    assert items[0]["price_source"] == "mayorista"
    assert items[0]["price_list_id"] == data["price_list"].id


def test_retail_quote_keeps_existing_price_and_product_discount(wholesale_database, app):
    import quotes

    data = wholesale_database
    with app.test_request_context("/presupuestos/nuevo"):
        login_user(data["employee"])
        try:
            items = quotes._parse_items(
                {"items": [{
                    "product_id": data["product"].id,
                    "description": data["product"].name,
                    "quantity": 2,
                    "unit_price": 100,
                    "discount": 10,
                }]},
                client=data["retail_client"],
            )
        finally:
            logout_user()

    assert items[0]["unit_price"] == Decimal("100")
    assert items[0]["discount"] == Decimal("10")
    assert items[0]["price_source"] == "minorista"
    assert items[0]["price_list_id"] is None


def test_admin_can_create_price_list_and_add_a_product_price(wholesale_database, app):
    data = wholesale_database
    client_id = int(data["wholesale_client"].id)
    browser = app.test_client()
    # Authenticate directly for this route test; the separate login flow has
    # its own tests, and this test focuses on administrator authorization.
    with browser.session_transaction() as session:
        session["_user_id"] = str(data["admin"].id)
        session["_fresh"] = True

    response = browser.post(
        "/precios-mayoristas/listas",
        data={"name": "Lista de prueba", "code": "lista-prueba"},
    )
    assert response.status_code in (301, 302)
    with app.app_context():
        price_list = PriceList.query.filter_by(
            company_id=data["company"].id, code="lista-prueba"
        ).first()
        assert price_list is not None
        list_id = price_list.id

    response = browser.post(
        f"/precios-mayoristas/listas/{list_id}/items",
        data={
            "product_id": str(data["product"].id),
            "min_quantity": "1",
            "unit_price": "75.50",
        },
    )
    assert response.status_code in (301, 302)
    with app.app_context():
        saved = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=list_id,
            product_id=data["product"].id,
            min_quantity=Decimal("1.000"),
        ).first()
        assert saved is not None
        assert Decimal(str(saved.unit_price)) == Decimal("75.50")

    management_page = browser.get(f"/precios-mayoristas/?list_id={list_id}")
    assert management_page.status_code == 200
    assert b"Tramos cargados" in management_page.data

    client_page = browser.get(f"/clientes/edit/{client_id}")
    assert client_page.status_code == 200
    assert b"Tipo de cliente" in client_page.data
    assert b"Lista de precios" in client_page.data


def test_wholesale_pricing_is_available_on_business_and_blocked_on_entrepreneur(wholesale_database, app):
    from services.ai_agent.usage_service import can_use_commercial_feature

    data = wholesale_database
    assert can_use_commercial_feature(data["company"], "wholesale_pricing").allowed is True

    browser = app.test_client()
    with browser.session_transaction() as session:
        session["_user_id"] = str(data["admin"].id)
        session["_fresh"] = True
    assert browser.get("/precios-mayoristas/").status_code == 200

    data["subscription"].plan_id = data["entrepreneur_plan"].id
    db.session.commit()
    assert can_use_commercial_feature(data["company"], "wholesale_pricing").allowed is False
    blocked = browser.get("/precios-mayoristas/")
    assert blocked.status_code == 302
    assert "precios-mayoristas" not in blocked.location

    resolver_result = resolve_product_price(
        company_id=data["company"].id,
        product=data["product"],
        quantity=12,
        client=data["wholesale_client"],
    )
    assert resolver_result["source"] == "minorista"
    assert resolver_result["reason"] == "plan_sin_mayoristas"
    assert resolver_result["price"] == Decimal("100.00")

    stale_quote_line = {
        "product": data["product"],
        "quantity": Decimal("12"),
        "price": Decimal("70.00"),
        "discount": Decimal("0.00"),
    }
    applied = apply_wholesale_prices(
        [stale_quote_line],
        company_id=data["company"].id,
        client=data["wholesale_client"],
    )
    assert applied[0]["price"] == Decimal("100.00")
    assert applied[0]["price_source"] == "minorista"
    assert product_price_catalog(
        company_id=data["company"].id,
        product_ids=[data["product"].id],
    ) == {}


def test_business_plan_catalog_includes_wholesale_pricing():
    business = next(plan for plan in PlanService.PLAN_CATALOG if plan["code"] == "business")
    features = {item.strip() for item in business["features_json"].split(",") if item.strip()}
    assert "wholesale_pricing" in features

def _make_wholesale_excel(rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or ["barcode", "product_name", "min_quantity", "unit_price", "active"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()



def test_wholesale_catalog_template_includes_all_active_products_and_existing_tiers(wholesale_database, app):
    data = wholesale_database
    second_product = Product(
        company_id=data["company"].id,
        barcode="WHOLESALE-002",
        name="Producto sin tarifa",
        price=Decimal("60.00"),
        cost_price=Decimal("30.00"),
        discount=Decimal("0.00"),
        stock=20,
        min_stock=1,
        active=True,
        sale_type="unidad",
        unit_measure="u",
    )
    db.session.add(second_product)
    db.session.commit()

    browser = app.test_client()
    with browser.session_transaction() as session:
        session["_user_id"] = str(data["admin"].id)
        session["_fresh"] = True

    response = browser.get(
        f"/precios-mayoristas/listas/{data['price_list'].id}/plantilla.xlsx"
    )
    assert response.status_code == 200
    assert "plantilla_precios_mayoristas_" in response.headers["Content-Disposition"]

    workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()

    assert rows[0] == ("barcode", "product_name", "min_quantity", "unit_price", "discount_percent", "active")
    assert len(rows) == 4  # header, two existing tiers and one product awaiting a price
    assert any(row[0] == "WHOLESALE-001" and row[2] == 1 and row[3] == 80 for row in rows[1:])
    assert any(row[0] == "WHOLESALE-001" and row[2] == 12 and row[3] == 70 for row in rows[1:])
    blank_price = next(row for row in rows[1:] if row[0] == "WHOLESALE-002")
    assert blank_price[1] == "Producto sin tarifa"
    assert blank_price[2] == 1
    assert blank_price[3] is None
    assert blank_price[4] is None
    assert blank_price[5] is True
    assert all(row[0] != "OTHER-001" for row in rows[1:])


def test_wholesale_import_skips_blank_price_rows_and_preserves_omitted_tiers(wholesale_database, app):
    data = wholesale_database
    list_id = data["price_list"].id
    product_id = data["product"].id
    browser = app.test_client()
    with browser.session_transaction() as session:
        session["_user_id"] = str(data["admin"].id)
        session["_fresh"] = True

    upload = _make_wholesale_excel([
        ["WHOLESALE-001", "Producto mayorista QA", 1, 77, True],
        ["WHOLESALE-001", "Producto mayorista QA", 12, None, True],
        ["WHOLESALE-001", "Producto mayorista QA", 48, None, True],
    ])
    response = browser.post(
        f"/precios-mayoristas/listas/{list_id}/importar.xlsx",
        data={"file": (BytesIO(upload), "precios.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code in (301, 302)

    with app.app_context():
        one_unit = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=list_id,
            product_id=product_id,
            min_quantity=Decimal("1.000"),
        ).first()
        twelve_units = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=list_id,
            product_id=product_id,
            min_quantity=Decimal("12.000"),
        ).first()
        forty_eight_units = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=list_id,
            product_id=product_id,
            min_quantity=Decimal("48.000"),
        ).first()
        assert one_unit is not None and Decimal(str(one_unit.unit_price)) == Decimal("77.00")
        assert twelve_units is not None and Decimal(str(twelve_units.unit_price)) == Decimal("70.00")
        assert forty_eight_units is None


def test_wholesale_price_excel_import_is_atomic_and_export_roundtrips(wholesale_database, app):
    data = wholesale_database
    with app.app_context():
        list_id = data["price_list"].id
        product_id = data["product"].id

    browser = app.test_client()
    with browser.session_transaction() as session:
        session["_user_id"] = str(data["admin"].id)
        session["_fresh"] = True

    invalid_upload = _make_wholesale_excel([
        ["WHOLESALE-001", "Producto mayorista QA", 1, 77, True],
        ["NO-EXISTE", "Producto no cargado", 1, 55, True],
    ])
    response = browser.post(
        f"/precios-mayoristas/listas/{list_id}/importar.xlsx",
        data={"file": (BytesIO(invalid_upload), "precios.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code in (301, 302)
    with app.app_context():
        original = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=list_id,
            product_id=product_id,
            min_quantity=Decimal("1.000"),
        ).first()
        assert Decimal(str(original.unit_price)) == Decimal("80.00")

    valid_upload = _make_wholesale_excel([
        ["WHOLESALE-001", "Producto mayorista QA", 1, 77, True],
        ["WHOLESALE-001", "Producto mayorista QA", 24, 65, True],
    ])
    response = browser.post(
        f"/precios-mayoristas/listas/{list_id}/importar.xlsx",
        data={"file": (BytesIO(valid_upload), "precios.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code in (301, 302)

    exported = browser.get(f"/precios-mayoristas/listas/{list_id}/exportar.xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.data), read_only=True, data_only=True)
    exported_rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    assert exported_rows[0] == ("barcode", "product_name", "min_quantity", "unit_price", "discount_percent", "active")
    assert any(row[0] == "WHOLESALE-001" and row[2] == 1 and row[3] == 77 for row in exported_rows[1:])
    assert any(row[0] == "WHOLESALE-001" and row[2] == 24 and row[3] == 65 for row in exported_rows[1:])

    with app.app_context():
        assert PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=list_id,
            product_id=product_id,
            min_quantity=Decimal("1.000"),
        ).first().unit_price == Decimal("77.00")



def test_saved_qr_price_snapshot_survives_current_product_price_change(wholesale_database):
    from types import SimpleNamespace

    data = wholesale_database
    product = SimpleNamespace(
        id=data["product"].id,
        company_id=data["company"].id,
        price=Decimal("125.00"),
    )
    lines = [{
        "product": product,
        "quantity": 12.0,
        "price": Decimal("125.00"),
        "discount": Decimal("0.00"),
    }]
    saved = [{
        "productId": product.id,
        "quantity": 12,
        "price": 70,
        "discount": 0,
        "price_source": "mayorista",
        "price_list_id": data["price_list"].id,
    }]

    result = apply_price_snapshot(lines, saved)

    assert result[0]["price"] == Decimal("70.00")
    assert result[0]["price_source"] == "mayorista"
    assert result[0]["price_list_id"] == data["price_list"].id


def test_saved_qr_price_snapshot_rejects_quantity_or_total_drift(wholesale_database):
    from types import SimpleNamespace

    data = wholesale_database
    product = SimpleNamespace(id=data["product"].id, company_id=data["company"].id)
    lines = [{"product": product, "quantity": 2.0, "price": Decimal("100"), "discount": Decimal("0")}]
    saved = [{
        "productId": product.id,
        "quantity": 3,
        "price": 80,
        "discount": 0,
        "price_source": "mayorista",
        "price_list_id": data["price_list"].id,
    }]
    with pytest.raises(ValueError, match="cantidad"):
        apply_price_snapshot(lines, saved)


def _create_wholesale_test_product(company_id, *, barcode, name, category, active=True, cost="20.00"):
    return Product(
        company_id=company_id,
        barcode=barcode,
        name=name,
        category=category,
        price=Decimal("100.00"),
        cost_price=Decimal(cost),
        discount=Decimal("0.00"),
        stock=10,
        min_stock=1,
        active=active,
        sale_type="unidad",
        unit_measure="u",
    )


def _wholesale_admin_client(app, user):
    browser = app.test_client()
    with browser.session_transaction() as session:
        session["_user_id"] = str(user.id)
        session["_fresh"] = True
    return browser


def test_wholesale_bulk_apply_to_all_active_products_is_tenant_scoped_and_preserves_other_tiers(wholesale_database, app):
    data = wholesale_database
    second_product = _create_wholesale_test_product(
        data["company"].id, barcode="WHOLESALE-002", name="Agua mineral", category="Bebidas"
    )
    inactive_product = _create_wholesale_test_product(
        data["company"].id, barcode="WHOLESALE-003", name="Producto inactivo", category="Bebidas", active=False
    )
    db.session.add_all([second_product, inactive_product])
    db.session.commit()
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/items",
        data={
            "bulk_mode": "all",
            "bulk_action": "replace_existing",
            "product_id": "",
            "category": "",
            "min_quantity": "1",
            "unit_price": "45",
        },
    )
    assert response.status_code in (301, 302)

    with app.app_context():
        original_base = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("1.000"),
        ).first()
        original_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("12.000"),
        ).first()
        second_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=second_product.id,
            min_quantity=Decimal("1.000"),
        ).first()
        inactive_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=inactive_product.id,
            min_quantity=Decimal("1.000"),
        ).first()
        cross_tenant_tiers = PriceListItem.query.filter_by(
            company_id=data["other_company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["other_product"].id,
        ).count()

        assert original_base is not None and Decimal(str(original_base.unit_price)) == Decimal("45.00")
        assert original_tier is not None and Decimal(str(original_tier.unit_price)) == Decimal("70.00")
        assert second_tier is not None and Decimal(str(second_tier.unit_price)) == Decimal("45.00")
        assert inactive_tier is None
        assert cross_tenant_tiers == 0


def test_wholesale_bulk_apply_by_category_only_changes_active_products_in_that_category(wholesale_database, app):
    data = wholesale_database
    second_beverage = _create_wholesale_test_product(
        data["company"].id, barcode="WHOLESALE-004", name="Jugo de naranja", category="Bebidas"
    )
    cleaning_product = _create_wholesale_test_product(
        data["company"].id, barcode="WHOLESALE-005", name="Detergente", category="Limpieza"
    )
    db.session.add_all([second_beverage, cleaning_product])
    db.session.commit()
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/items",
        data={
            "bulk_mode": "category",
            "bulk_action": "replace_existing",
            "product_id": "",
            "category": "Bebidas",
            "min_quantity": "12",
            "unit_price": "66.50",
        },
    )
    assert response.status_code in (301, 302)

    with app.app_context():
        beverage_existing = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("12.000"),
        ).first()
        beverage_new = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=second_beverage.id,
            min_quantity=Decimal("12.000"),
        ).first()
        cleaning_row = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=cleaning_product.id,
            min_quantity=Decimal("12.000"),
        ).first()
        other_company_row = PriceListItem.query.filter_by(
            price_list_id=data["price_list"].id,
            product_id=data["other_product"].id,
            min_quantity=Decimal("12.000"),
        ).first()

        assert beverage_existing is not None and Decimal(str(beverage_existing.unit_price)) == Decimal("66.50")
        assert beverage_new is not None and Decimal(str(beverage_new.unit_price)) == Decimal("66.50")
        assert cleaning_row is None
        assert other_company_row is None



def test_wholesale_bulk_fill_missing_preserves_existing_prices(wholesale_database, app):
    data = wholesale_database
    missing_product = _create_wholesale_test_product(
        data["company"].id,
        barcode="WHOLESALE-MISSING-001",
        name="Producto sin tramo base",
        category="Bebidas",
    )
    disabled_tier_product = _create_wholesale_test_product(
        data["company"].id,
        barcode="WHOLESALE-INACTIVE-TIER",
        name="Producto con tramo desactivado",
        category="Limpieza",
    )
    db.session.add_all([missing_product, disabled_tier_product])
    db.session.flush()
    db.session.add(PriceListItem(
        company_id=data["company"].id,
        price_list_id=data["price_list"].id,
        product_id=disabled_tier_product.id,
        min_quantity=Decimal("1.000"),
        unit_price=Decimal("55.00"),
        active=False,
    ))
    db.session.commit()
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/items",
        data={
            "bulk_mode": "all",
            "bulk_action": "fill_missing",
            "product_id": "",
            "category": "",
            "min_quantity": "1",
            "unit_price": "45",
        },
    )
    assert response.status_code in (301, 302)

    with app.app_context():
        existing_price = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("1.000"),
        ).first()
        added_price = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=missing_product.id,
            min_quantity=Decimal("1.000"),
        ).first()
        other_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("12.000"),
        ).first()

        assert existing_price is not None
        assert Decimal(str(existing_price.unit_price)) == Decimal("80.00")
        assert added_price is not None
        assert Decimal(str(added_price.unit_price)) == Decimal("45.00")
        assert other_tier is not None
        assert Decimal(str(other_tier.unit_price)) == Decimal("70.00")
        reactivated = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=disabled_tier_product.id,
            min_quantity=Decimal("1.000"),
        ).first()
        assert reactivated is not None
        assert reactivated.active is True
        assert Decimal(str(reactivated.unit_price)) == Decimal("45.00")


def test_wholesale_bulk_fill_missing_with_no_missing_tiers_does_not_require_price(wholesale_database, app):
    data = wholesale_database
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/items",
        data={
            "bulk_mode": "all",
            "bulk_action": "fill_missing",
            "product_id": "",
            "category": "",
            "min_quantity": "1",
            "unit_price": "",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "No se modificó ningún precio".encode() in response.data

    with app.app_context():
        existing_price = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("1.000"),
        ).first()
        assert existing_price is not None
        assert Decimal(str(existing_price.unit_price)) == Decimal("80.00")


def test_wholesale_bulk_preview_counts_existing_and_missing_tiers(wholesale_database, app):
    data = wholesale_database
    missing_product = _create_wholesale_test_product(
        data["company"].id,
        barcode="WHOLESALE-PREVIEW-001",
        name="Producto sin precio",
        category="Bebidas",
    )
    db.session.add(missing_product)
    db.session.commit()
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.get(
        f"/precios-mayoristas/listas/{data['price_list'].id}/previsualizar-masivo"
        "?bulk_mode=all&min_quantity=1"
    )
    assert response.status_code == 200
    payload = response.get_json()
    assert payload == {
        "total": 2,
        "already_priced": 1,
        "missing": 1,
        "min_quantity": "1.000",
    }


def test_wholesale_product_search_and_category_selector_are_company_scoped(wholesale_database, app):
    data = wholesale_database
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.get("/precios-mayoristas/productos/buscar?q=WHOLESALE")
    assert response.status_code == 200
    payload = response.get_json()
    assert payload is not None
    assert [item["id"] for item in payload["products"]] == [data["product"].id]
    assert all("OTHER-001" not in item["barcode"] for item in payload["products"])

    page = browser.get(f"/precios-mayoristas/?list_id={data['price_list'].id}")
    assert page.status_code == 200
    assert b"Todos los productos activos" in page.data
    assert b"Una categor" in page.data
    assert b"Bebidas" in page.data
    assert b"Descuento % sobre precio minorista" in page.data


def test_wholesale_bulk_percentage_discount_calculates_each_product_and_tracks_retail_price(wholesale_database, app):
    data = wholesale_database
    second_product = _create_wholesale_test_product(
        data["company"].id,
        barcode="WHOLESALE-PCT-002",
        name="Producto de precio minorista doble",
        category="Bebidas",
        cost="40.00",
    )
    second_product.price = Decimal("200.00")
    db.session.add(second_product)
    db.session.commit()
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/items",
        data={
            "bulk_mode": "all",
            "bulk_action": "fill_missing",
            "pricing_mode": "discount_percent",
            "product_id": "",
            "category": "",
            "min_quantity": "3",
            "unit_price": "",
            "discount_percent": "15",
        },
    )
    assert response.status_code in (301, 302)

    with app.app_context():
        product_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("3.000"),
        ).first()
        second_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=second_product.id,
            min_quantity=Decimal("3.000"),
        ).first()
        assert product_tier is not None
        assert Decimal(str(product_tier.discount_percent)) == Decimal("15.00")
        assert Decimal(str(product_tier.unit_price)) == Decimal("85.00")
        assert second_tier is not None
        assert Decimal(str(second_tier.discount_percent)) == Decimal("15.00")
        assert Decimal(str(second_tier.unit_price)) == Decimal("170.00")

        # Existing quantities remain as before and percentage tiers are dynamic.
        base_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("1.000"),
        ).first()
        assert Decimal(str(base_tier.unit_price)) == Decimal("80.00")
        second_product.price = Decimal("400.00")
        db.session.commit()

        result = resolve_product_price(
            company_id=data["company"].id,
            product=second_product,
            quantity=3,
            client=data["wholesale_client"],
        )
        assert result["price"] == Decimal("340.00")
        assert result["discount_percent"] == Decimal("15.00")

        catalog = product_price_catalog(
            company_id=data["company"].id,
            product_ids=[second_product.id],
        )
        catalog_tiers = catalog[second_product.id][str(data["price_list"].id)]
        quantity_three = next(tier for tier in catalog_tiers if tier["min_quantity"] == 3.0)
        assert quantity_three["price"] == 340.0
        assert quantity_three["discount_percent"] == 15.0


def test_wholesale_bulk_percentage_discount_replaces_existing_tier_only_when_requested(wholesale_database, app):
    data = wholesale_database
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/items",
        data={
            "bulk_mode": "all",
            "bulk_action": "replace_existing",
            "pricing_mode": "discount_percent",
            "product_id": "",
            "category": "",
            "min_quantity": "1",
            "unit_price": "",
            "discount_percent": "20",
        },
    )
    assert response.status_code in (301, 302)

    with app.app_context():
        tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("1.000"),
        ).first()
        other_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("12.000"),
        ).first()
        assert tier is not None
        assert Decimal(str(tier.discount_percent)) == Decimal("20.00")
        assert Decimal(str(tier.unit_price)) == Decimal("80.00")
        assert other_tier is not None
        assert other_tier.discount_percent is None
        assert Decimal(str(other_tier.unit_price)) == Decimal("70.00")


def test_wholesale_discount_percent_excel_roundtrips_and_accepts_legacy_templates(wholesale_database, app):
    data = wholesale_database
    browser = _wholesale_admin_client(app, data["admin"])
    headers = ["barcode", "product_name", "min_quantity", "unit_price", "discount_percent", "active"]
    upload = _make_wholesale_excel(
        [["WHOLESALE-001", "Producto mayorista QA", 3, None, 15, True]],
        headers=headers,
    )
    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/importar.xlsx",
        data={"file": (BytesIO(upload), "descuentos.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code in (301, 302)

    with app.app_context():
        tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("3.000"),
        ).first()
        assert tier is not None
        assert Decimal(str(tier.discount_percent)) == Decimal("15.00")
        assert Decimal(str(tier.unit_price)) == Decimal("85.00")

    exported = browser.get(f"/precios-mayoristas/listas/{data['price_list'].id}/exportar.xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.data), read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    pct_row = next(row for row in rows[1:] if row[0] == "WHOLESALE-001" and row[2] == 3)
    assert pct_row[3] is None
    assert pct_row[4] == 15
    assert pct_row[5] is True

    # Older exported sheets have no discount_percent column and still import as fixed prices.
    legacy_upload = _make_wholesale_excel(
        [["WHOLESALE-001", "Producto mayorista QA", 24, 64, True]]
    )
    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/importar.xlsx",
        data={"file": (BytesIO(legacy_upload), "legacy.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code in (301, 302)
    with app.app_context():
        legacy_tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("24.000"),
        ).first()
        assert legacy_tier is not None
        assert legacy_tier.discount_percent is None
        assert Decimal(str(legacy_tier.unit_price)) == Decimal("64.00")


def test_wholesale_discount_percentage_requires_value_below_100(wholesale_database, app):
    data = wholesale_database
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/items",
        data={
            "bulk_mode": "all",
            "bulk_action": "fill_missing",
            "pricing_mode": "discount_percent",
            "product_id": "",
            "category": "",
            "min_quantity": "3",
            "unit_price": "",
            "discount_percent": "100",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "mayor a 0% y menor a 100%".encode() in response.data

    with app.app_context():
        created = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("3.000"),
        ).first()
        assert created is None



def test_single_product_can_use_percentage_discount_tier(wholesale_database, app):
    data = wholesale_database
    browser = _wholesale_admin_client(app, data["admin"])

    response = browser.post(
        f"/precios-mayoristas/listas/{data['price_list'].id}/items",
        data={
            "bulk_mode": "single",
            "pricing_mode": "discount_percent",
            "product_id": str(data["product"].id),
            "min_quantity": "3",
            "unit_price": "",
            "discount_percent": "10",
        },
    )
    assert response.status_code in (301, 302)

    with app.app_context():
        tier = PriceListItem.query.filter_by(
            company_id=data["company"].id,
            price_list_id=data["price_list"].id,
            product_id=data["product"].id,
            min_quantity=Decimal("3.000"),
        ).first()
        assert tier is not None
        assert Decimal(str(tier.discount_percent)) == Decimal("10.00")
        assert Decimal(str(tier.unit_price)) == Decimal("90.00")

        data["product"].price = Decimal("200.00")
        db.session.commit()
        resolution = resolve_product_price(
            company_id=data["company"].id,
            product=data["product"],
            quantity=3,
            client=data["wholesale_client"],
        )
        assert resolution["price"] == Decimal("180.00")

