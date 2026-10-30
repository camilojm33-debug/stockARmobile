from io import BytesIO
from urllib.parse import urlsplit

import pytest
from flask import url_for
from openpyxl import load_workbook

from app import Company, Product, ProductModification, User, db
from stockarmobile.helpers.strings import is_safe_relative_redirect


@pytest.fixture
def authenticated_tenant_client(app):
    with app.app_context():
        company = Company(name="Empresa seguridad QA", active=True)
        db.session.add(company)
        db.session.flush()
        user = User(
            username="security_qa_admin",
            email="security-qa@example.test",
            role="admin",
            active=True,
            company_id=company.id,
        )
        user.set_password("test-password-123")
        db.session.add(user)
        db.session.commit()

    client = app.test_client()
    response = client.post(
        "/auth/login",
        data={"username": "security_qa_admin", "password": "test-password-123"},
    )
    assert response.status_code in (301, 302)
    return client


@pytest.mark.parametrize(
    "target",
    [
        "https://attacker.example/path",
        "//attacker.example/path",
        "///attacker.example/path",
        r"/\attacker.example/path",
        "/%2f%2fattacker.example/path",
        "/%252f%252fattacker.example/path",
        "/%5c%5cattacker.example/path",
        "javascript:alert(1)",
    ],
)
def test_redirect_validator_rejects_external_and_ambiguous_destinations(target):
    assert not is_safe_relative_redirect(target)


@pytest.mark.parametrize(
    "target",
    [
        "/productos",
        "/clientes?page=2",
        "/dashboard/stats#ventas",
    ],
)
def test_redirect_validator_allows_internal_paths(target):
    assert is_safe_relative_redirect(target)


def test_login_rejects_triple_slash_next(authenticated_tenant_client, app):
    response = authenticated_tenant_client.post(
        "/auth/login?next=///attacker.example/path",
        data={"username": "security_qa_admin", "password": "test-password-123"},
    )

    with app.test_request_context():
        expected_location = url_for("dashboard.index")
    assert response.status_code in (301, 302)
    assert response.headers["Location"] == expected_location
    assert not urlsplit(response.headers["Location"]).netloc


def test_login_preserves_a_valid_internal_next(authenticated_tenant_client):
    response = authenticated_tenant_client.post(
        "/auth/login?next=/productos",
        data={"username": "security_qa_admin", "password": "test-password-123"},
    )

    assert response.status_code in (301, 302)
    assert response.headers["Location"] == "/productos"


def test_tour_completion_rejects_triple_slash_next(authenticated_tenant_client, app):
    response = authenticated_tenant_client.post(
        "/dashboard/tour/complete",
        data={"next": "///attacker.example/path"},
    )

    with app.test_request_context():
        expected_location = url_for("dashboard.index")
    assert response.status_code in (301, 302)
    assert response.headers["Location"] == expected_location


def test_tour_completion_preserves_a_valid_internal_next(authenticated_tenant_client):
    response = authenticated_tenant_client.post(
        "/dashboard/tour/complete",
        data={"next": "/productos"},
    )

    assert response.status_code in (301, 302)
    assert response.headers["Location"] == "/productos"


def test_product_xlsx_export_keeps_user_text_as_text(authenticated_tenant_client, app):
    unsafe_values = [
        "=1+1",
        "+1+1",
        "-1+1",
        "@SUM(1,1)",
        "\t=1+1",
        "\r=1+1",
        "\n=1+1",
        " =1+1",
    ]
    with app.app_context():
        company = Company.query.filter_by(name="Empresa seguridad QA").one()
        product = Product(
            company_id=company.id,
            barcode=unsafe_values[0],
            name=unsafe_values[1],
            category=unsafe_values[2],
            brand=unsafe_values[3],
            supplier=unsafe_values[4],
            sale_type=unsafe_values[5],
            unit_measure=unsafe_values[6],
            active=True,
        )
        db.session.add(product)
        db.session.add(
            Product(
                company_id=company.id,
                barcode="SKU-LEGITIMO",
                name="Producto legítimo ñ á",
                category="Bebidas",
                brand="Marca",
                supplier="Proveedor",
                sale_type="unidad",
                unit_measure="u",
                active=True,
            )
        )
        db.session.add(
            Product(
                company_id=company.id,
                barcode="SKU-PREFIJO-ESPACIO",
                name=unsafe_values[7],
                category="Bebidas",
                brand="Marca",
                supplier="Proveedor",
                sale_type="unidad",
                unit_measure="u",
                active=True,
            )
        )
        db.session.commit()

    response = authenticated_tenant_client.get("/productos/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), data_only=False)
    sheet = workbook.active
    rows_by_barcode = {
        row[0].value: row
        for row in sheet.iter_rows(min_row=2)
    }

    unsafe_row = rows_by_barcode[unsafe_values[0]]
    for column, expected in zip((1, 2, 3, 4, 5, 10, 11), unsafe_values):
        cell = unsafe_row[column - 1]
        assert cell.value == expected.replace("\r", "\n")
        assert cell.data_type == "s"

    legitimate_row = rows_by_barcode["SKU-LEGITIMO"]
    assert legitimate_row[1].value == "Producto legítimo ñ á"
    assert legitimate_row[1].data_type == "s"
    assert legitimate_row[5].value == 0
    space_prefixed_row = rows_by_barcode["SKU-PREFIJO-ESPACIO"]
    assert space_prefixed_row[1].value == unsafe_values[7]
    assert space_prefixed_row[1].data_type == "s"
    workbook.close()


def test_product_kardex_xlsx_keeps_user_detail_as_text(authenticated_tenant_client, app):
    malicious_detail = '=HYPERLINK("https://attacker.example","open")'
    with app.app_context():
        company = Company.query.filter_by(name="Empresa seguridad QA").one()
        product = Product(
            company_id=company.id,
            barcode="SKU-KARDEX",
            name="Producto kardex",
            active=True,
        )
        db.session.add(product)
        db.session.flush()
        db.session.add(
            ProductModification(
                product_id=product.id,
                company_id=company.id,
                action="nota",
                detail=malicious_detail,
            )
        )
        db.session.commit()
        product_id = product.id

    response = authenticated_tenant_client.get(
        f"/productos/{product_id}/kardex/export.xlsx"
    )
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), data_only=False)
    cell = workbook.active["C2"]
    assert cell.value == malicious_detail
    assert cell.data_type == "s"
    workbook.close()


@pytest.mark.parametrize(
    "target",
    [
        "https://attacker.example/path",
        "//attacker.example/path",
        "///attacker.example/path",
        r"/\\attacker.example/path",
        "/%2f%2fattacker.example/path",
        "/%252f%252fattacker.example/path",
        "/%5c%5cattacker.example/path",
        "/%25255c%25255cattacker.example/path",
        "javascript:alert(1)",
    ],
)
@pytest.mark.parametrize("source", ["query", "form"])
def test_saas_redirect_back_rejects_external_and_ambiguous_destinations(app, target, source):
    from saas import _redirect_back

    if source == "query":
        request_context = app.test_request_context("/", query_string={"next": target})
    else:
        request_context = app.test_request_context("/", method="POST", data={"next": target})

    with request_context:
        response = _redirect_back(default_endpoint="dashboard.index")
        assert response.status_code == 302
        assert response.headers["Location"] == url_for("dashboard.index")


@pytest.mark.parametrize("source", ["query", "form"])
def test_saas_redirect_back_preserves_valid_internal_destinations(app, source):
    from saas import _redirect_back

    target = "/clientes?page=2"
    if source == "query":
        request_context = app.test_request_context("/", query_string={"next": target})
    else:
        request_context = app.test_request_context("/", method="POST", data={"next": target})

    with request_context:
        response = _redirect_back(default_endpoint="dashboard.index")
        assert response.status_code == 302
        assert response.headers["Location"] == target
