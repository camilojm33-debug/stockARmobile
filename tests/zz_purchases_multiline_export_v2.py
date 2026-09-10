import io

from openpyxl import Workbook, load_workbook

from app import Company, Product, PurchaseItem, PurchaseOrder, Supplier, User, db


def _login(client, user):
    with client.session_transaction() as session:
        session["_user_id"] = str(user.id)
        session["_fresh"] = True


def _setup():
    company = Company(name="Empresa Compras Test", active=True)
    db.session.add(company)
    db.session.flush()
    user = User(username="compras_admin", email="compras-admin@test.local", password_hash="test-password-hash", role="admin", active=True, company_id=company.id)
    supplier = Supplier(company_id=company.id, name="Proveedor Test", active=True)
    p1 = Product(barcode="CMP-A", name="Producto A", price=150, cost_price=100, stock=0, min_stock=1, active=True, company_id=company.id)
    p2 = Product(barcode="CMP-B", name="Producto B", price=300, cost_price=200, stock=0, min_stock=1, active=True, company_id=company.id)
    db.session.add_all([user, supplier, p1, p2])
    db.session.commit()
    return company, user, supplier, p1, p2


def test_multiline_purchase_post_creates_one_order_and_two_items(app):
    client = app.test_client()
    with app.app_context():
        company, user, supplier, p1, p2 = _setup()
        _login(client, user)
        response = client.post("/compras/", data={
            "company_id": str(company.id), "supplier_id": str(supplier.id),
            "product_id[]": [str(p1.id), str(p2.id)], "quantity[]": ["2", "3"],
            "unit_cost[]": ["25", "50"], "note": "multi",
        })
        assert response.status_code in {302, 303}
        purchase = PurchaseOrder.query.filter_by(company_id=company.id, supplier_id=supplier.id).one()
        assert len(purchase.items) == 2
        assert float(purchase.total_amount) == 200.0
        assert float(db.session.get(Product, p1.id).stock) == 2.0
        assert float(db.session.get(Product, p2.id).stock) == 3.0


def test_supplier_purchase_excel_contains_real_data(app):
    client = app.test_client()
    with app.app_context():
        company, user, supplier, p1, p2 = _setup()
        purchase = PurchaseOrder(supplier_id=supplier.id, company_id=company.id, status="recibida", subtotal=200, total_amount=200)
        db.session.add(purchase)
        db.session.flush()
        db.session.add_all([
            PurchaseItem(purchase_order_id=purchase.id, product_id=p1.id, quantity=2, unit_cost=25),
            PurchaseItem(purchase_order_id=purchase.id, product_id=p2.id, quantity=3, unit_cost=50),
        ])
        db.session.commit()
        _login(client, user)
        response = client.get(f"/compras/proveedores/{supplier.id}/compras/export.xlsx?company_id={company.id}")
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
        rows = list(workbook["Compras"].iter_rows(values_only=True))
        workbook.close()
        assert any(row[3] == "Producto A" and row[4] == 2.0 for row in rows[1:])
        assert any(row[3] == "Producto B" and row[4] == 3.0 for row in rows[1:])


def test_purchase_edit_replaces_multiple_items_and_reconciles_stock(app):
    client = app.test_client()
    with app.app_context():
        company, user, supplier, p1, p2 = _setup()
        purchase = PurchaseOrder(supplier_id=supplier.id, company_id=company.id, status="recibida", subtotal=200, total_amount=200)
        db.session.add(purchase)
        db.session.flush()
        db.session.add_all([
            PurchaseItem(purchase_order_id=purchase.id, product_id=p1.id, quantity=2, unit_cost=25),
            PurchaseItem(purchase_order_id=purchase.id, product_id=p2.id, quantity=3, unit_cost=50),
        ])
        p1.stock, p2.stock = 2, 3
        db.session.commit()
        _login(client, user)
        response = client.post(f"/compras/proveedores/{supplier.id}/compras/{purchase.id}/editar", data={
            "company_id": str(company.id), "product_id[]": [str(p1.id), str(p2.id)],
            "quantity[]": ["4", "1"], "unit_cost[]": ["30", "55"],
            "status": "recibida", "note": "editada",
        })
        assert response.status_code in {302, 303}
        refreshed = db.session.get(PurchaseOrder, purchase.id)
        assert len(refreshed.items) == 2
        assert float(refreshed.total_amount) == 175.0
        assert float(db.session.get(Product, p1.id).stock) == 4.0
        assert float(db.session.get(Product, p2.id).stock) == 1.0


def test_import_excel_creates_one_multiproduct_purchase(app):
    client = app.test_client()
    with app.app_context():
        company, user, supplier, p1, p2 = _setup()
        _login(client, user)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Producto", "Cantidad", "Costo unitario"])
        sheet.append([p1.name, 2, 25])
        sheet.append([p2.name, 3, 50])
        payload = io.BytesIO()
        workbook.save(payload)
        payload.seek(0)
        response = client.post(
            f"/compras/proveedores/{supplier.id}/compras/import.xlsx?company_id={company.id}",
            data={"file": (payload, "compras.xlsx")}, content_type="multipart/form-data",
        )
        assert response.status_code in {302, 303}
        purchase = PurchaseOrder.query.filter_by(company_id=company.id, supplier_id=supplier.id).one()
        assert len(purchase.items) == 2
        assert float(purchase.total_amount) == 200.0
